import os, openpyxl, shutil
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

src = r"C:\Users\burha\OneDrive\Desktop\1 project\portfolio\career\data\applications.xlsx"
dst = r"C:\Users\burha\OneDrive\Desktop\1 project\portfolio\career\data\Vertex Job Tracker.xlsx"

# rename: load & save as new name (aman untuk xlsx)
if os.path.exists(src):
    wb = openpyxl.load_workbook(src)
    wb.save(dst)
    os.remove(src)
    print("renamed ->", dst)
else:
    # kalau sudah tidak ada, buat baru
    wb = openpyxl.Workbook()

ws = wb["Applications"] if "Applications" in wb.sheetnames else wb.active
ws.title = "Applications"
hdr = ["Tanggal_Kirim","Perusahaan","Posisi","Sumber_JD","Status","Skor_AG","Link_CV_PDF","Followup_Terakhir","Catatan","Laporan_Evaluasi"]
# tambah kolom ke-10 jika belum
if ws.max_column < 10:
    ws.cell(row=1, column=10).value = "Laporan_Evaluasi"
    hdr = [c.value for c in ws[1]]
fill = PatternFill("solid", fgColor="26794A")
hf = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin,right=thin,top=thin,bottom=thin)
for c in range(1, ws.max_column+1):
    cell=ws.cell(row=1,column=c); cell.fill=fill; cell.font=hf
    cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=border

# isi/perbarui baris ke-2 dengan data BPDLH lengkap
report_link = "reports/001-BPDLH-TenagaPendukung-2026-08-26.md"
cv_link = "data/CV_BurhanuddienRobbani_ATS.pdf"
row2 = ["2026-08-26","BPDLH / Kementerian Kehutanan","Tenaga Pendukung (REDD+ GCF Output 2)","procurement.bpdlh.id/file_tor/GCF/01_Tenaga_pendukung.pdf","Draft (belum kirim)","B","","","Gap: English terbatas; relevansi sektor tinggi. Perlu cover letter + bukti tulisan", report_link]
for i,v in enumerate(row2,1):
    ws.cell(row=2,column=i).value=v
for r in range(2, ws.max_row+1):
    for c in range(1, ws.max_column+1):
        cell=ws.cell(row=r,column=c); cell.border=border
        if r==2: cell.alignment=Alignment(wrap_text=True, vertical="top")
widths=[13,26,30,34,16,9,30,16,40,40]
for i,w in enumerate(widths,1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes="A2"
wb.save(dst)
print("UPDATED:", dst, os.path.getsize(dst), "bytes")
print("rows:", ws.max_row, "cols:", ws.max_column)
