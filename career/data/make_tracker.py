import os, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

csvp = r"C:\Users\burha\OneDrive\Desktop\1 project\portfolio\career\data\applications.csv"
wb = openpyxl.Workbook(); ws = wb.active; ws.title="Applications"
hdr = ["Tanggal_Kirim","Perusahaan","Posisi","Sumber_JD","Status","Skor_AG","Link_CV_PDF","Followup_Terakhir","Catatan"]
ws.append(hdr)
fill = PatternFill("solid", fgColor="26794A")
hf = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="CCCCCC")
border = Border(left=thin,right=thin,top=thin,bottom=thin)
for c in range(1,len(hdr)+1):
    cell=ws.cell(row=1,column=c); cell.fill=fill; cell.font=hf
    cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=border
ws.append(["","BPDLH / GCF","Tenaga Pendukung (lihat TOR)","procurement.bpdlh.id","Belum dievaluasi","","","","JD terblokir Cloudflare - tunggu PDF dari Robbani"])
for r in range(2, ws.max_row+1):
    for c in range(1,len(hdr)+1):
        ws.cell(row=r,column=c).border=border
widths=[14,26,24,22,16,10,30,16,36]
for i,w in enumerate(widths,1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes="A2"
xlsx = r"C:\Users\burha\OneDrive\Desktop\1 project\portfolio\career\data\applications.xlsx"
wb.save(xlsx)
print("XLSX tracker dibuat:", xlsx, os.path.getsize(xlsx), "bytes")
if os.path.exists(csvp): os.remove(csvp); print("CSV lama dihapus")